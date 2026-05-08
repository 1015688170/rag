from __future__ import annotations

import hashlib
import json
import re
import shutil
import uuid
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, BinaryIO

from fastapi import UploadFile
from sqlalchemy.orm import Session

from app.core.database import UPLOAD_DIR
from app.core.permissions import list_to_json, parse_string_list
from app.models.knowledge import Document, IngestTask
from app.schemas.chat import EmbeddingModel
from app.services.embedding_service import EmbeddingService
from app.services.search_service import SearchService


@dataclass
class DocumentSection:
    section_title: str
    section_path: str
    source_type: str
    content: str
    page_start: int | None = None
    page_end: int | None = None


@dataclass
class DocumentChunk:
    chunk_index: int
    section_title: str
    section_path: str
    source_type: str
    content: str
    page_start: int | None = None
    page_end: int | None = None


class DocumentIngestService:
    supported_extensions = {".json", ".md", ".txt", ".docx", ".pdf"}
    max_file_size = 20 * 1024 * 1024
    chunk_size = 1000
    chunk_overlap = 150

    def __init__(self, embedding_service: EmbeddingService, search_service: SearchService) -> None:
        self.embedding_service = embedding_service
        self.search_service = search_service

    async def ingest(
        self,
        db: Session,
        file: UploadFile,
        index_name: str,
        embedding_model: EmbeddingModel,
        visibility: str = "public",
        owner_id: str | None = None,
        allowed_departments: list[str] | None = None,
        allowed_roles: list[str] | None = None,
    ) -> dict:
        filename = self.validate_file(file)
        document_id = uuid.uuid4().hex
        task_id = uuid.uuid4().hex
        saved_path, file_size = self.save_file(file.file, filename, document_id)
        file_hash = self.calculate_file_hash(saved_path)
        duplicate = self.check_duplicate(db, file_hash)
        if duplicate and duplicate.status == "success":
            task = self._create_task(db, task_id, duplicate.id, filename, "success", "duplicate", duplicate.chunk_count)
            shutil.rmtree(saved_path.parent, ignore_errors=True)
            return self._response(duplicate, task, status="already_exists", index_name=index_name, embedding_model=embedding_model)

        if duplicate:
            shutil.rmtree(saved_path.parent, ignore_errors=True)
            document = duplicate
            task = self._create_task(db, task_id, document.id, document.filename, "pending", "saved", 0)
        else:
            document, task = self._create_document_and_task(
                db,
                document_id,
                task_id,
                filename,
                saved_path,
                file_hash,
                file_size,
                visibility,
                owner_id,
                allowed_departments or [],
                allowed_roles or [],
            )

        try:
            self.update_document_status(db, document, "parsing")
            self.update_task_status(db, task, "parsing", "parsing")
            sections = self.parse_document(Path(document.filepath), document.file_type)

            self.update_document_status(db, document, "chunking")
            self.update_task_status(db, task, "chunking", "chunking")
            chunks = self.split_chunks(sections)
            if not chunks:
                raise ValueError("Uploaded document has no extractable text.")

            self.update_document_status(db, document, "embedding")
            self.update_task_status(db, task, "embedding", "embedding", chunk_count=len(chunks))
            indexed_chunks = self.embed_chunks(document, chunks, embedding_model)

            self.update_document_status(db, document, "indexing", chunk_count=len(indexed_chunks))
            self.update_task_status(db, task, "indexing", "indexing", chunk_count=len(indexed_chunks))
            uploaded_count = self.index_chunks_to_azure_search(index_name, indexed_chunks)
            if uploaded_count != len(indexed_chunks):
                raise RuntimeError(f"Only {uploaded_count}/{len(indexed_chunks)} chunks were indexed.")

            self.update_document_status(db, document, "success", chunk_count=len(indexed_chunks))
            self.update_task_status(db, task, "success", "success", chunk_count=len(indexed_chunks))
        except Exception as exc:
            message = str(exc)
            self.update_document_status(db, document, "failed", error_message=message)
            self.update_task_status(db, task, "failed", "failed", error_message=message)
            raise

        db.refresh(document)
        db.refresh(task)
        return self._response(document, task, index_name=index_name, embedding_model=embedding_model)

    def validate_file(self, file: UploadFile) -> str:
        filename = self._safe_filename(file.filename or "")
        extension = Path(filename).suffix.lower()
        if not filename or extension not in self.supported_extensions:
            supported = ", ".join(sorted(self.supported_extensions))
            raise ValueError(f"Unsupported file type. Supported: {supported}.")
        return filename

    def save_file(self, source: BinaryIO, filename: str, document_id: str) -> tuple[Path, int]:
        target_dir = UPLOAD_DIR / document_id
        target_dir.mkdir(parents=True, exist_ok=True)
        target_path = target_dir / filename
        size = 0
        with target_path.open("wb") as output:
            while True:
                chunk = source.read(1024 * 1024)
                if not chunk:
                    break
                size += len(chunk)
                if size > self.max_file_size:
                    shutil.rmtree(target_dir, ignore_errors=True)
                    raise ValueError("File too large. Maximum size is 20MB.")
                output.write(chunk)
        if size == 0:
            shutil.rmtree(target_dir, ignore_errors=True)
            raise ValueError("Uploaded file is empty.")
        return target_path, size

    def calculate_file_hash(self, filepath: Path) -> str:
        digest = hashlib.sha256()
        with filepath.open("rb") as handle:
            for block in iter(lambda: handle.read(1024 * 1024), b""):
                digest.update(block)
        return digest.hexdigest()

    def check_duplicate(self, db: Session, file_hash: str) -> Document | None:
        return db.query(Document).filter(Document.file_hash == file_hash).first()

    def parse_document(self, filepath: Path, file_type: str) -> list[DocumentSection]:
        suffix = f".{file_type.lower().lstrip('.')}"
        if suffix == ".json":
            return self._parse_json(filepath)
        if suffix == ".md":
            return self._parse_markdown(filepath)
        if suffix == ".txt":
            return self._parse_plain_text(filepath, "txt")
        if suffix == ".docx":
            return self._parse_docx(filepath)
        if suffix == ".pdf":
            return self._parse_pdf(filepath)
        raise ValueError(f"Unsupported file type: {suffix}.")

    def split_chunks(self, sections: list[DocumentSection]) -> list[DocumentChunk]:
        chunks: list[DocumentChunk] = []
        for section in sections:
            text = self._normalize_text(section.content)
            if not text:
                continue
            for part in self._split_section_content(text):
                chunks.append(
                    DocumentChunk(
                        chunk_index=len(chunks),
                        section_title=section.section_title,
                        section_path=section.section_path,
                        source_type=section.source_type,
                        content=part,
                        page_start=section.page_start,
                        page_end=section.page_end,
                    )
                )
        return chunks

    def embed_chunks(
        self,
        document: Document,
        chunks: list[DocumentChunk],
        embedding_model: EmbeddingModel,
    ) -> list[dict[str, Any]]:
        created_at = datetime.utcnow().isoformat(timespec="seconds") + "Z"
        indexed_chunks: list[dict[str, Any]] = []
        for chunk in chunks:
            chunk_id = f"{document.id}_{chunk.chunk_index}"
            indexed_chunks.append(
                {
                    "id": chunk_id,
                    "doc_id": document.id,
                    "chunk_id": chunk_id,
                    "chunk_index": chunk.chunk_index,
                    "filename": document.filename,
                    "filepath": document.filepath,
                    "section_title": chunk.section_title,
                    "section_path": chunk.section_path,
                    "source_type": chunk.source_type,
                    "content": chunk.content,
                    "content_vector": self.embedding_service.embed(chunk.content, embedding_model),
                    "page_start": chunk.page_start,
                    "page_end": chunk.page_end,
                    "created_at": created_at,
                    "file_hash": document.file_hash,
                    "visibility": document.visibility,
                    "owner_id": document.owner_id,
                    "allowed_departments": parse_string_list(document.allowed_departments),
                    "allowed_roles": parse_string_list(document.allowed_roles),
                }
            )
        return indexed_chunks

    def index_chunks_to_azure_search(self, index_name: str, chunks: list[dict[str, Any]]) -> int:
        return self.search_service.index_chunks(index_name=index_name, chunks=chunks)

    def update_document_status(
        self,
        db: Session,
        document: Document,
        status: str,
        chunk_count: int | None = None,
        error_message: str | None = None,
    ) -> None:
        document.status = status
        document.error_message = error_message
        document.updated_at = datetime.utcnow()
        if chunk_count is not None:
            document.chunk_count = chunk_count
        db.commit()

    def update_task_status(
        self,
        db: Session,
        task: IngestTask,
        status: str,
        stage: str | None = None,
        chunk_count: int | None = None,
        error_message: str | None = None,
    ) -> None:
        task.status = status
        task.stage = stage
        task.error_message = error_message
        task.updated_at = datetime.utcnow()
        if chunk_count is not None:
            task.chunk_count = chunk_count
        db.commit()

    def _parse_markdown(self, filepath: Path) -> list[DocumentSection]:
        text = filepath.read_text(encoding="utf-8-sig")
        sections: list[DocumentSection] = []
        current_title = "$"
        current_path = "$"
        heading_stack: list[str] = []
        current_lines: list[str] = []
        for line in text.splitlines():
            heading = re.match(r"^(#{1,6})\s+(.+?)\s*$", line)
            if heading and current_lines:
                sections.append(DocumentSection(current_title, current_path, "md", "\n".join(current_lines)))
                current_lines = []
            if heading:
                level = len(heading.group(1))
                current_title = heading.group(2).strip()
                heading_stack = heading_stack[: level - 1]
                heading_stack.append(current_title)
                current_path = "/".join(heading_stack)
            current_lines.append(line)
        if current_lines:
            sections.append(DocumentSection(current_title, current_path, "md", "\n".join(current_lines)))
        return sections

    def _parse_plain_text(self, filepath: Path, source_type: str) -> list[DocumentSection]:
        text = filepath.read_text(encoding="utf-8-sig")
        paragraphs = [part.strip() for part in re.split(r"\n\s*\n", text) if part.strip()]
        return self._sections_from_paragraphs(paragraphs, source_type)

    def _parse_docx(self, filepath: Path) -> list[DocumentSection]:
        try:
            from docx import Document as DocxDocument
        except ImportError as exc:
            raise RuntimeError("DOCX upload requires the python-docx package.") from exc

        doc = DocxDocument(str(filepath))
        sections: list[DocumentSection] = []
        heading_stack: list[str] = []
        current_title = "$"
        current_path = "$"
        current_lines: list[str] = []
        has_heading = False
        plain_paragraphs: list[str] = []

        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style_name = paragraph.style.name if paragraph.style is not None else ""
            heading = re.match(r"Heading\s+(\d+)", style_name, flags=re.IGNORECASE)
            if heading:
                has_heading = True
                if current_lines:
                    sections.append(DocumentSection(current_title, current_path, "docx", "\n\n".join(current_lines)))
                    current_lines = []
                level = int(heading.group(1))
                current_title = text
                heading_stack = heading_stack[: level - 1]
                heading_stack.append(text)
                current_path = "/".join(heading_stack)
                current_lines.append(text)
            else:
                current_lines.append(text)
                plain_paragraphs.append(text)

        if has_heading:
            if current_lines:
                sections.append(DocumentSection(current_title, current_path, "docx", "\n\n".join(current_lines)))
            return sections
        return self._sections_from_paragraphs(plain_paragraphs, "docx")

    def _parse_pdf(self, filepath: Path) -> list[DocumentSection]:
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("PDF upload requires the pypdf package.") from exc

        reader = PdfReader(str(filepath))
        paragraphs: list[tuple[str, int]] = []
        for index, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            # TODO: filter repeated headers, footers, and standalone page numbers.
            for paragraph in re.split(r"\n\s*\n", text):
                normalized = self._normalize_text(paragraph)
                if normalized:
                    paragraphs.append((normalized, index + 1))
        if not paragraphs:
            raise ValueError("PDF text extraction returned no text. Scanned PDFs/OCR are not supported in this version.")
        return self._sections_from_paged_paragraphs(paragraphs, "pdf")

    def _parse_json(self, filepath: Path) -> list[DocumentSection]:
        try:
            data = json.loads(filepath.read_text(encoding="utf-8-sig"))
        except json.JSONDecodeError as exc:
            raise ValueError(f"JSON parse error: {exc.msg} at line {exc.lineno}, column {exc.colno}.") from exc

        sections = self._json_sections(data, "$")
        if not sections:
            sections = [DocumentSection("$", "$", "json", json.dumps(data, ensure_ascii=False, indent=2))]
        return sections

    def _json_sections(self, value: Any, path: str) -> list[DocumentSection]:
        serialized = json.dumps(value, ensure_ascii=False, indent=2)
        if len(serialized) <= self.chunk_size or not isinstance(value, (dict, list)) or not value:
            return [DocumentSection(path, path, "json", serialized)]

        if isinstance(value, dict):
            sections: list[DocumentSection] = []
            for key, child in value.items():
                child_path = f"{path}.{key}" if path != "$" else f"$.{key}"
                sections.extend(self._json_sections(child, child_path))
            return sections

        sections = []
        for index, item in enumerate(value):
            item_path = f"{path}[{index}]"
            sections.extend(self._json_sections(item, item_path))
        return sections

    def _sections_from_paragraphs(self, paragraphs: list[str], source_type: str) -> list[DocumentSection]:
        sections: list[DocumentSection] = []
        current: list[str] = []
        for paragraph in paragraphs:
            text = self._normalize_text(paragraph)
            if not text:
                continue
            if current and len("\n\n".join(current)) + len(text) + 2 > self.chunk_size:
                title = f"paragraph[{len(sections)}]"
                sections.append(DocumentSection(title, title, source_type, "\n\n".join(current)))
                current = []
            current.append(text)
        if current:
            title = f"paragraph[{len(sections)}]"
            sections.append(DocumentSection(title, title, source_type, "\n\n".join(current)))
        return sections

    def _sections_from_paged_paragraphs(
        self,
        paragraphs: list[tuple[str, int]],
        source_type: str,
    ) -> list[DocumentSection]:
        sections: list[DocumentSection] = []
        current: list[str] = []
        page_start: int | None = None
        page_end: int | None = None

        for paragraph, page_number in paragraphs:
            if current and len("\n\n".join(current)) + len(paragraph) + 2 > self.chunk_size:
                title = f"page[{page_start}]" if page_start == page_end else f"page[{page_start}-{page_end}]"
                sections.append(
                    DocumentSection(title, title, source_type, "\n\n".join(current), page_start, page_end)
                )
                current = []
                page_start = None
                page_end = None

            current.append(paragraph)
            page_start = page_number if page_start is None else page_start
            page_end = page_number

        if current:
            title = f"page[{page_start}]" if page_start == page_end else f"page[{page_start}-{page_end}]"
            sections.append(DocumentSection(title, title, source_type, "\n\n".join(current), page_start, page_end))
        return sections

    def _split_section_content(self, text: str) -> list[str]:
        if len(text) <= self.chunk_size:
            return [text]

        units = self._semantic_units(text)
        chunks: list[str] = []
        current = ""
        for unit in units:
            if not unit:
                continue
            if len(unit) > self.chunk_size:
                if current:
                    chunks.append(current.strip())
                    current = ""
                chunks.extend(self._char_windows(unit))
                continue

            candidate = unit if not current else f"{current}\n\n{unit}"
            if len(candidate) <= self.chunk_size:
                current = candidate
                continue

            if current:
                chunks.append(current.strip())
                overlap = self._overlap_suffix(current)
                current = f"{overlap}\n\n{unit}" if overlap else unit
                if len(current) > self.chunk_size:
                    chunks.extend(self._char_windows(current))
                    current = ""
            else:
                current = unit

        if current.strip():
            chunks.append(current.strip())
        return [chunk for chunk in chunks if chunk.strip()]

    def _semantic_units(self, text: str) -> list[str]:
        units = [part.strip() for part in re.split(r"\n\s*\n", text) if part.strip()]
        if len(units) >= 2 and re.fullmatch(r"#{1,6}\s+.+", units[0]):
            units = [f"{units[0]}\n\n{units[1]}", *units[2:]]
        return units or [text]

    def _char_windows(self, text: str) -> list[str]:
        chunks: list[str] = []
        start = 0
        while start < len(text):
            end = min(start + self.chunk_size, len(text))
            chunk = text[start:end].strip()
            if chunk:
                chunks.append(chunk)
            if end == len(text):
                break
            start = max(end - self.chunk_overlap, start + 1)
        return chunks

    def _overlap_suffix(self, text: str) -> str:
        suffix = text[-self.chunk_overlap :].strip()
        first_break = suffix.find("\n")
        if first_break > 0:
            suffix = suffix[first_break:].strip()
        return suffix

    def _normalize_text(self, text: str) -> str:
        return re.sub(r"\n{3,}", "\n\n", text).strip()

    def _create_document_and_task(
        self,
        db: Session,
        document_id: str,
        task_id: str,
        filename: str,
        saved_path: Path,
        file_hash: str,
        file_size: int,
        visibility: str,
        owner_id: str | None,
        allowed_departments: list[str],
        allowed_roles: list[str],
    ) -> tuple[Document, IngestTask]:
        now = datetime.utcnow()
        document = Document(
            id=document_id,
            filename=filename,
            filepath=str(saved_path),
            file_hash=file_hash,
            file_type=Path(filename).suffix.lower().lstrip("."),
            file_size=file_size,
            chunk_count=0,
            status="pending",
            error_message=None,
            visibility=visibility,
            owner_id=owner_id,
            allowed_departments=list_to_json(allowed_departments),
            allowed_roles=list_to_json(allowed_roles),
            created_at=now,
            updated_at=now,
        )
        task = IngestTask(
            id=task_id,
            document_id=document_id,
            filename=filename,
            status="pending",
            stage="saved",
            error_message=None,
            chunk_count=0,
            created_at=now,
            updated_at=now,
        )
        db.add(document)
        db.add(task)
        db.commit()
        db.refresh(document)
        db.refresh(task)
        return document, task

    def _create_task(
        self,
        db: Session,
        task_id: str,
        document_id: str,
        filename: str,
        status: str,
        stage: str,
        chunk_count: int,
    ) -> IngestTask:
        now = datetime.utcnow()
        task = IngestTask(
            id=task_id,
            document_id=document_id,
            filename=filename,
            status=status,
            stage=stage,
            error_message=None,
            chunk_count=chunk_count,
            created_at=now,
            updated_at=now,
        )
        db.add(task)
        db.commit()
        db.refresh(task)
        return task

    def _response(
        self,
        document: Document,
        task: IngestTask,
        status: str | None = None,
        index_name: str | None = None,
        embedding_model: EmbeddingModel | None = None,
    ) -> dict:
        return {
            "document_id": document.id,
            "task_id": task.id,
            "filename": document.filename,
            "file_hash": document.file_hash,
            "chunk_count": document.chunk_count,
            "status": status or document.status,
            "index_name": index_name,
            "embedding_model": embedding_model,
            "visibility": document.visibility,
            "owner_id": document.owner_id,
            "allowed_departments": parse_string_list(document.allowed_departments),
            "allowed_roles": parse_string_list(document.allowed_roles),
        }

    def _safe_filename(self, filename: str) -> str:
        name = Path(filename).name.strip().replace("\\", "_").replace("/", "_")
        return re.sub(r"[^A-Za-z0-9._ -]", "_", name)
